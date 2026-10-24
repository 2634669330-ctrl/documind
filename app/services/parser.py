"""文档解析：把 PDF / Word / Markdown / TXT 统一成「页 + 文本」的结构。

返回的 Page 列表是后面分块和引用溯源的基础：
- PDF 用真实页码（引用能精确到「第 42 页」）；
- Word / Markdown / TXT 没有可靠的物理分页，按约 2500 字符切成逻辑页，
  这样引用里仍然能给出「第 N 页」这种可读定位。
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

# 无分页格式的“逻辑页”大小（字符）
_PSEUDO_PAGE_CHARS = 2500


class UnsupportedFormatError(ValueError):
    """扩展名不在支持列表里。"""


class ParseError(RuntimeError):
    """文件损坏、加密、或提取不到文本。"""


@dataclass(frozen=True)
class Page:
    no: int
    text: str


@dataclass
class ParsedDocument:
    filename: str
    pages: list[Page] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


def parse_file(path: Path, filename: str | None = None) -> ParsedDocument:
    """根据扩展名选择解析器。path 已落盘，filename 只用于展示与元数据。"""
    display_name = filename or path.name
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        pages = _parse_pdf(path)
    elif suffix == ".docx":
        pages = _parse_docx(path)
    elif suffix in (".md", ".markdown", ".txt"):
        pages = _parse_text_like(path)
    else:
        raise UnsupportedFormatError(f"暂不支持的文件类型：{suffix}")

    pages = [p for p in pages if p.text.strip()]
    if not pages:
        raise ParseError("未提取到任何文本，可能是扫描件或空文件")

    logger.info("解析完成 filename=%s pages=%d chars=%d", display_name, len(pages), sum(len(p.text) for p in pages))
    return ParsedDocument(filename=display_name, pages=pages)


# --------------------------------------------------------------------------- PDF
def _parse_pdf(path: Path) -> list[Page]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少依赖 pypdf，请先 pip install -r requirements.txt") from exc

    try:
        reader = PdfReader(str(path))
        # 空密码能解开的就解，解不开直接报错，提示用户处理
        if reader.is_encrypted and not reader.decrypt(""):
            raise ParseError("PDF 已加密，需要先解除密码保护")
        pages: list[Page] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(Page(no=idx, text=_clean(text)))
        return pages
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"PDF 解析失败：{exc}") from exc


# -------------------------------------------------------------------------- DOCX
def _parse_docx(path: Path) -> list[Page]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少依赖 python-docx，请先 pip install -r requirements.txt") from exc

    try:
        document = DocxDocument(str(path))
        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        # 表格里的文字也要收进来，很多规范是表格形式
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return _split_pseudo_pages("\n\n".join(blocks))
    except Exception as exc:
        raise ParseError(f"Word 解析失败：{exc}") from exc


# ------------------------------------------------------------------ MD / TXT
def _parse_text_like(path: Path) -> list[Page]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ParseError("无法识别文件编码（尝试过 utf-8 / gb18030）")
    return _split_pseudo_pages(_clean(text))


# --------------------------------------------------------------------------- 工具
def _clean(text: str) -> str:
    """去掉 PDF 常见的零宽字符与多余空行。"""
    text = text.replace("\u200b", "").replace("\xa0", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_pseudo_pages(text: str, size: int = _PSEUDO_PAGE_CHARS) -> list[Page]:
    text = text.strip()
    if not text:
        return []
    pages: list[Page] = []
    for idx, start in enumerate(range(0, len(text), size), start=1):
        piece = text[start:start + size].strip()
        if piece:
            pages.append(Page(no=idx, text=piece))
    return pages
